"""Reproducible synthetic XLSX audit benchmark; no Excel or arbitrary eval needed."""
from pathlib import Path
from collections import Counter, defaultdict
import ast, csv, hashlib, json, platform, re, time, zipfile
import numpy as np
import openpyxl, sklearn
from openpyxl.styles import PatternFill
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries, column_index_from_string
from sklearn.mixture import GaussianMixture
from sklearn.ensemble import IsolationForest
from sklearn.preprocessing import StandardScaler
from sklearn.metrics import average_precision_score, roc_auc_score
from threadpoolctl import threadpool_limits

ROOT = Path(__file__).resolve().parent
SEED = 20261003
ROWS = 180
COLORS = ['FFF2CC','FFF2CC','D9EAD3','FFF2CC','CFE2F3','D9D2E9']
SCENARIOS = ['semantic', 'noisy', 'decorative', 'reversed']

def workbook(seed, scenario, defects):
    rng = np.random.default_rng(seed)
    style_rng = np.random.default_rng(seed+10000)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Ledger'
    ws.append(['Quantity','Price','Gross','Adjustment','Net','Running net'])
    for r in range(2, ROWS+2):
        quantity = float(rng.normal(10 if rng.random()<.7 else 35, 2))
        if rng.random()<.02: quantity *= -1  # legitimate returns
        price = float(rng.lognormal(2, .25))
        adjustment = float(rng.normal(0, 3))
        ws.append([quantity, price, f'=A{r}*B{r}', adjustment,
                   f'=SUM(C{r},D{r})' if rng.random()<.04 else f'=C{r}+D{r}',
                   f'=E{r}' if r==2 else f'=F{r-1}+E{r}'])
        for c in range(1,7):
            color = COLORS[c-1]
            if scenario == 'noisy' and style_rng.random()<.2: color = style_rng.choice(COLORS)
            if scenario == 'decorative': color = COLORS[(r+c)%6]
            if scenario == 'reversed': color = {'D9EAD3':'CFE2F3','CFE2F3':'D9EAD3'}.get(color,color)
            ws.cell(r,c).fill = PatternFill('solid', fgColor=color)
            ws.cell(r,c).number_format = '0.00'
    summary = wb.create_sheet('Summary'); summary['A1']='Net total'
    summary['B1']=f'=SUM(Ledger!E2:E{ROWS+1})'
    labels = {}
    if defects:
        for r, kind in zip(rng.choice(np.arange(3,ROWS+2),24,replace=False),
                           ['wrong_operator','wrong_reference','hardcode','numeric_scale']*6):
            r=int(r)
            if kind=='wrong_operator': address=f'C{r}'; ws[address]=f'=A{r}+B{r}'
            elif kind=='wrong_reference': address=f'E{r}'; ws[address]=f'=C{r-1}+D{r}'
            elif kind=='hardcode':
                address=f'E{r}'; ws[address]=float(ws[f'A{r}'].value*ws[f'B{r}'].value+ws[f'D{r}'].value)
            else: address=f'A{r}'; ws[address]=float(ws[address].value*30)
            labels['Ledger!'+address]=kind
    return wb, labels

def refs(formula, sheet):
    out=[]
    for tok in Tokenizer(formula).items:
        if tok.type=='OPERAND' and tok.subtype=='RANGE':
            value=tok.value
            if '!' in value: target, value=value.rsplit('!',1); target=target.strip("'")
            else: target=sheet
            lo_c,lo_r,hi_c,hi_r=range_boundaries(value.replace('$',''))
            for r in range(lo_r,hi_r+1):
                for c in range(lo_c,hi_c+1): out.append((target,r,c))
    return out

def evaluate(wb):
    cache={}; visiting=set()
    def value(sheet,r,c):
        key=(sheet,r,c)
        if key in cache: return cache[key]
        if key in visiting: raise ValueError('Cycle detected')
        visiting.add(key); raw=wb[sheet].cell(r,c).value
        if isinstance(raw,(float,int)): result=float(raw)
        elif isinstance(raw,str) and raw.startswith('='):
            bits=[]
            for tok in Tokenizer(raw).items:
                if tok.type=='OPERAND' and tok.subtype=='RANGE':
                    resolved=refs('='+tok.value,sheet)
                    bits.append(','.join(repr(value(*x)) for x in resolved))
                else: bits.append(tok.value)
            tree=ast.parse(''.join(bits),mode='eval')
            def calc(node):
                if isinstance(node,ast.Expression): return calc(node.body)
                if isinstance(node,ast.Constant) and isinstance(node.value,(int,float)): return node.value
                if isinstance(node,ast.UnaryOp) and isinstance(node.op,ast.USub): return -calc(node.operand)
                if isinstance(node,ast.BinOp):
                    a,b=calc(node.left),calc(node.right)
                    if isinstance(node.op,ast.Add): return a+b
                    if isinstance(node.op,ast.Sub): return a-b
                    if isinstance(node.op,ast.Mult): return a*b
                    if isinstance(node.op,ast.Div): return a/b
                if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and node.func.id=='SUM':
                    return sum(calc(x) for x in node.args)
                raise ValueError('Unsupported formula syntax')
            result=float(calc(tree))
        else: result=0.0
        visiting.remove(key); cache[key]=result; return result
    for sheet in wb:
        for row in sheet:
            for cell in row: value(sheet.title,cell.row,cell.column)
    return cache

def extract(path, labels):
    wb=openpyxl.load_workbook(path,data_only=False); values=evaluate(wb)
    edges={}; fanout=Counter()
    for sheet in wb:
        for row in sheet:
            for cell in row:
                key=(sheet.title,cell.row,cell.column)
                edges[key]=refs(cell.value,sheet.title) if cell.data_type=='f' else []
                fanout.update(edges[key])
    cells=[]
    for row in wb['Ledger'].iter_rows(min_row=2):
        for cell in row:
            key=('Ledger',cell.row,cell.column); dependencies=edges[key]
            f=cell.value if cell.data_type=='f' else ''
            offsets=[(r-cell.row,c-cell.column) for s,r,c in dependencies]
            numeric=values[key]
            # Continuous embedding of values and structural counts; color is categorical separately.
            feature=[np.sign(numeric)*np.log1p(abs(numeric)),bool(f),len(dependencies),
                     max([abs(r) for r,c in offsets] or [0]),
                     sum(abs(c) for r,c in offsets),f.count('*'),f.count('+'),
                     'SUM(' in f,np.log1p(fanout[key])]
            signature='literal' if not f else re.sub(r'\$?([A-Z]+)\$?(\d+)',
                lambda m:f'R{int(m[2])-cell.row}C{column_index_from_string(m[1])-cell.column}',f)
            address=f'Ledger!{cell.coordinate}'
            cells.append({'address':address,'x':feature,'signature':signature,
                          'style':cell.fill.fgColor.rgb,'col':cell.column,
                          'label':int(address in labels),'kind':labels.get(address,'clean'),
                          'fanout':fanout[key]})
    return cells, sum(len(x) for x in edges.values())

def fit(cells):
    x=np.array([c['x'] for c in cells]); scaler=StandardScaler().fit(x); z=scaler.transform(x)
    candidates=[]
    for k in [2,4,6,8]:
        model=GaussianMixture(k,covariance_type='diag',reg_covar=1e-3,n_init=3,random_state=SEED).fit(z)
        candidates.append((model.bic(z),model))
    bic,gmm=min(candidates,key=lambda item:item[0])
    iso=IsolationForest(n_estimators=100,random_state=SEED,n_jobs=1).fit(z)
    counts=Counter(c['signature'] for c in cells); styles=defaultdict(Counter); columns=defaultdict(Counter)
    for c in cells: styles[c['style']][c['signature']]+=1; columns[c['col']][c['signature']]+=1
    vocab=len(counts)+1
    def categorical(c, color):
        distribution=styles[c['style']] if color else counts
        return -np.log((distribution[c['signature']]+1)/(sum(distribution.values())+vocab))
    train_nll=-gmm.score_samples(z)
    train_cat=np.array([categorical(c,False) for c in cells]); train_color=np.array([categorical(c,True) for c in cells])
    return locals()

def scores(cells, model):
    z=model['scaler'].transform(np.array([c['x'] for c in cells]))
    nll=-model['gmm'].score_samples(z)
    cat=np.array([model['categorical'](c,False) for c in cells])
    color=np.array([model['categorical'](c,True) for c in cells])
    def rank(test, train): return np.searchsorted(np.sort(train),test,side='right')/len(train)
    n=rank(nll,model['train_nll']); s=rank(cat,model['train_cat']); f=rank(color,model['train_color'])
    rules=np.array([int(c['signature']!=model['columns'][c['col']].most_common(1)[0][0]) for c in cells],dtype=float)
    return {'column_mode':rules,'global_gmm':nll,'isolation_forest':-model['iso'].score_samples(z),
            'structure_gmm':.5*n+.5*s,'format_gmm':.5*n+.5*f}

def metrics(cells, score, seed):
    y=np.array([c['label'] for c in cells]); rng=np.random.default_rng(seed)
    # Random tie breaking is shared across methods and independent of labels.
    order=np.lexsort((rng.random(len(y)),-score)); k=int(np.ceil(.05*len(y))); hits=y[order[:k]].sum()
    result={'ap':float(average_precision_score(y,score)),'auroc':float(roc_auc_score(y,score)),
            'precision_at_5pct':float(hits/k),'recall_at_5pct':float(hits/y.sum()),'budget':k}
    for kind in ['wrong_operator','wrong_reference','hardcode','numeric_scale']:
        truth=np.array([c['kind']==kind for c in cells]); result['recall_'+kind]=float(truth[order[:k]].sum()/truth.sum())
    return result

def main():
    start=time.perf_counter(); out=ROOT/'results'; out.mkdir(exist_ok=True)
    data=out/'workbooks'; data.mkdir(exist_ok=True); training=[]; hashes={}; runs=[]; labels_all={}
    for i in range(8):
        wb,labels=workbook(SEED+i,'semantic',False); path=data/f'train-{i:02d}.xlsx'; wb.save(path)
        cells,edges=extract(path,labels); training.extend(cells)
    with threadpool_limits(limits=1): model=fit(training)
    for scenario in SCENARIOS:
        for i in range(12):
            seed=SEED+100+i
            wb,labels=workbook(seed,scenario,True); path=data/f'{scenario}-{i:02d}.xlsx'; wb.save(path)
            labels_all[path.name]=labels
            cells,edges=extract(path,labels)
            for method,score in scores(cells,model).items():
                runs.append({'scenario':scenario,'replicate':i,'method':method,'cells':len(cells),'defects':len(labels),
                             'edges':edges,**metrics(cells,score,seed)})
    keys=list(runs[0]);
    with (out/'runs.csv').open('w') as handle:
        writer=csv.DictWriter(handle,fieldnames=keys); writer.writeheader(); writer.writerows(runs)
    summary=[]; rng=np.random.default_rng(SEED)
    for scenario in SCENARIOS:
        for method in scores(cells,model):
            group=[r for r in runs if r['scenario']==scenario and r['method']==method]
            row={'scenario':scenario,'method':method}
            for metric in ['ap','auroc','precision_at_5pct','recall_at_5pct','recall_wrong_operator','recall_wrong_reference','recall_hardcode','recall_numeric_scale']:
                v=np.array([g[metric] for g in group]); boot=v[rng.integers(0,len(v),size=(2000,len(v)))].mean(axis=1)
                row[metric]=float(v.mean()); row[metric+'_ci']=[float(x) for x in np.quantile(boot,[.025,.975])]
            summary.append(row)
    (out/'summary.json').write_text(json.dumps(summary,indent=2)); (out/'labels.json').write_text(json.dumps(labels_all,indent=2))
    for path in data.iterdir(): hashes[path.name]=hashlib.sha256(path.read_bytes()).hexdigest()
    manifest={'seed':SEED,'training_workbooks':8,'test_workbooks':48,'rows_per_ledger':ROWS,
              'scored_cells_per_workbook':ROWS*6,'defects_per_test_workbook':24,'selected_components':model['gmm'].n_components,
              'bic_candidates':{str(m.n_components):float(b) for b,m in model['candidates']},
              'converged':bool(model['gmm'].converged_),'python':platform.python_version(),
              'numpy':np.__version__,'sklearn':sklearn.__version__,'openpyxl':openpyxl.__version__,
              'elapsed_seconds':time.perf_counter()-start,'sha256':hashes}
    (out/'manifest.json').write_text(json.dumps(manifest,indent=2))
    with zipfile.ZipFile(out/'workbooks.zip','w',zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(data.iterdir()): archive.write(path,'workbooks/'+path.name)
        archive.write(out/'labels.json','labels.json')
    for row in summary: print(row['scenario'],row['method'],round(row['ap'],3),round(row['recall_at_5pct'],3))
    print('Manifest:',json.dumps({k:v for k,v in manifest.items() if k!='sha256'}))

if __name__=='__main__': main()
