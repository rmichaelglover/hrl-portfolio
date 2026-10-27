"""Generate only fictional lighting-quote data and known review examples."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
wb=Workbook();ws=wb.active;ws.title='Lighting Quote'
ws.append(['Item','Quantity','Unit price','Line amount','Notes'])
for r in range(2,21):ws.append([f'Fixture {r-1:02d}',r%5+1,round(24.5+r*3.2,2),f'=B{r}*C{r}','Fictional sample'])
ws['D6']='=B6+C6';ws['D11']=248.5;ws['E11']='Intentional negotiated fixed amount: legitimate exception'
ws['D16']='=B16*#REF!'
ws['C18']='#DIV/0!';ws['C18'].data_type='e'
ws['B19']=-2;ws['E19']='Legitimate return: a negative quantity is not itself flagged'
ws['D22']='=SUM(D2:D20)';ws['C22']='Total'
for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='245E4B')
for r in range(2,23):
 for col in ['C','D']:ws[f'{col}{r}'].number_format='#,##0.00'
for col,width in [('A',23),('B',14),('C',16),('D',18),('E',65)]:ws.column_dimensions[col].width=width
ws.freeze_panes='A2'
notes=wb.create_sheet('Read me');notes.append(['Fictional workbook — no company or customer data'])
for text in ['D6: inconsistent formula (addition instead of multiplication).','D11: intentional fixed-price override; mark Intentional.','D16: broken reference.','C18: saved Excel error.','B19: negative return is legitimate and should not be flagged.','D22: subtotal is outside the repeated row block and should not be flagged.','Formulas are not recalculated by the review app.']:notes.append([text])
wb.save(Path(__file__).resolve().parent/'sample-lighting-quote.xlsx')
